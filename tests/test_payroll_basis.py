"""Расчёт зарплаты по большему из записей и отмеченных посещений."""
import sqlite3
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import db, lessons
from app.admin_check import compare

DAY = date(2026, 10, 15)

@pytest.fixture
def conn():
    c = sqlite3.connect(':memory:')
    c.row_factory = sqlite3.Row
    c.executescript(db.SCHEMA)
    yield c
    c.close()


def seed(c, gid, visits, booked, cancelled=0):
    c.execute('INSERT INTO groups (id,style,teacher1,updated_at) VALUES (?,?,?,?)',
              (gid, f'Группа {gid}', 'Педагог', '2026-10-17'))
    c.execute('INSERT INTO sessions VALUES (?,?,?,?,?,?,?,?)',
              (gid, DAY.isoformat(), gid, 1, 1200, 1260, 1, cancelled))
    if visits:
        c.execute('INSERT INTO visits_agg VALUES (?,?,?,?,?,?,?)',
                  (DAY.isoformat(), gid, 1200, 1, 'groupAccount', gid, visits))
    c.execute('INSERT INTO bookings VALUES (?,?,?)', (gid, DAY.isoformat(), booked))


@pytest.mark.parametrize('visits,booked,total,pay',[(3,4,4,1000),(9,7,9,1500),(6,6,6,1000),(0,4,4,1000)])
def test_basis_and_payment(conn,visits,booked,total,pay):
    seed(conn,1,visits,booked)
    row=lessons.lesson_rows(conn,DAY,DAY).iloc[0]
    assert (row['Человек'],row['Записались'],row['Для зарплаты'],row['Оплата, ₽']) == (visits,booked,total,pay)


def test_joint_basis_is_max_after_summing(conn):
    seed(conn,1,3,0)
    seed(conn,2,0,4)
    rows=lessons.lesson_rows(conn,DAY,DAY)
    assert len(rows)==1
    r=rows.iloc[0]
    assert r['joint']
    assert (r['Человек'],r['Записались'],r['Для зарплаты'],r['Оплата, ₽']) == (3,4,4,1000)


def test_cancelled_and_empty_are_not_paid(conn):
    seed(conn,1,3,4,cancelled=1)
    seed(conn,2,0,0)
    rows=lessons.lesson_rows(conn,DAY,DAY)
    assert not rows['counts'].any()
    assert rows['Для зарплаты'].isna().all()
    assert rows['Оплата, ₽'].isna().all()


def test_admin_comparison_uses_payroll_basis(conn):
    seed(conn,1,3,4)
    result,money,_,_=compare(conn,'15.10\nПедагог - 4',date(2026,10,17))
    assert result.iloc[0]['Итог']=='✅'
    assert result.iloc[0]['Бот']==4
    assert result.iloc[0]['Отмечены']==3
    assert money['Разница'].sum()==0


def test_text_and_excel_show_both_counts(conn,tmp_path,monkeypatch):
    seed(conn,1,3,4)
    rows=lessons.lesson_rows(conn,DAY,DAY)
    text=lessons.build_text(rows,DAY,DAY,[])
    assert 'для зарплаты: 4; записались: 4; отмечены: 3 · 1 000 ₽' in text
    monkeypatch.setattr(lessons,'REPORTS_DIR',tmp_path)
    path=lessons.build_excel(rows,DAY,DAY)
    wb=load_workbook(path,data_only=True)
    ws=wb['Занятия']; columns={c.value:c.column for c in ws[1]}
    for name,value in [('Для зарплаты',4),('Записались',4),('Отмечены',3),('Оплата, ₽',1000)]:
        assert ws.cell(2,columns[name]).value==value
    ws=wb['По педагогам']; columns={c.value:c.column for c in ws[1]}
    assert ws.cell(2,columns['Для зарплаты всего']).value==4
    assert ws.cell(2,columns['Отмечены всего']).value==3
    wb.close()
