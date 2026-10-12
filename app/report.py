"""Еженедельный отчёт о посещаемости: текст (для Telegram) и Excel-файл."""

from datetime import date, timedelta

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app import config
from app.attendance import loaded_weeks, recompute_all, week_totals
from app.metrics import (STATUS_CANCELLED, STATUS_DROP, STATUS_GROW, STATUS_NEW, STATUS_OK,
                         STATUS_WATCH)

DAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]
REPORTS_DIR = config.PROJECT_DIR / "reports"
STATUS_ORDER = [STATUS_DROP, STATUS_WATCH, STATUS_GROW, STATUS_CANCELLED, STATUS_NEW, STATUS_OK]
TOP_GROWING = 3


def hhmm(minutes) -> str:
    if minutes is None or pd.isna(minutes):
        return "без времени"
    minutes = int(minutes)
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def group_labels(conn, monday: date | None = None) -> dict[int, str]:
    """Группа → «Afro · Аня К. · пн 19:00, ср 19:00».

    Время занятий — за неделю отчёта, а если в ней занятий не было — за последнюю неделю группы.
    """
    labels = {}
    for g in conn.execute("SELECT id, style, teacher1 FROM groups"):
        until = (monday + timedelta(days=6)).isoformat() if monday else "9999-12-31"
        slots = conn.execute(
            """SELECT DISTINCT minutes_begin, CAST(strftime('%w', date) AS INTEGER) AS wd
               FROM sessions WHERE group_id = ? AND regular = 1
                 AND date >= date((SELECT MAX(date) FROM sessions WHERE group_id = ? AND date <= ?), '-6 days')
                 AND date <= ?
               ORDER BY (wd + 6) % 7, minutes_begin""",
            (g["id"], g["id"], until, until)).fetchall()
        when = ", ".join(f"{DAYS[(s['wd'] + 6) % 7]} {hhmm(s['minutes_begin'])}" for s in slots)
        labels[g["id"]] = " · ".join(x for x in (g["style"], g["teacher1"], when) if x)
    return labels


def last_full_week(today: date | None = None) -> date:
    """Понедельник прошлой недели."""
    today = today or date.today()
    return today - timedelta(days=today.weekday() + 7)


def week_of(d: date) -> date:
    return d - timedelta(days=d.weekday())


def fmt_num(x) -> str:
    return "—" if x is None or pd.isna(x) else f"{x:.1f}".replace(".", ",")


def fmt_pct(x, signed=False) -> str:
    if x is None or pd.isna(x):
        return "—"
    return f"{x:+.0%}" if signed else f"{x:.0%}"


# ---------- данные отчёта ----------

def week_stats(conn, monday: date) -> pd.DataFrame:
    df = pd.read_sql_query("SELECT * FROM weekly_group_stats WHERE week_start = ?", conn,
                           params=(monday.isoformat(),))
    df["Группа"] = df["group_id"].map(group_labels(conn, monday))
    df["order"] = df["status"].map({s: i for i, s in enumerate(STATUS_ORDER)})
    return df.sort_values(["order", "change_pct", "Группа"]).reset_index(drop=True)


def week_changes(conn, monday: date) -> pd.DataFrame:
    """Отмены и замены недели по занятиям из расписания (отмены удалённых шаблонов не показываем)."""
    df = pd.read_sql_query(
        """SELECT c.date, c.entity, c.source_teacher, c.teacher, s.group_id, s.minutes_begin
           FROM changes c JOIN sessions s ON s.schedule_id = c.schedule_id AND s.date = c.date
           WHERE c.date BETWEEN ? AND ? ORDER BY c.date, s.minutes_begin""",
        conn, params=(monday.isoformat(), (monday + timedelta(days=6)).isoformat()))
    labels = group_labels(conn, monday)
    df["Группа"] = df["group_id"].map(lambda g: " · ".join(labels[g].split(" · ")[:2]))  # направление · преподаватель
    return df


# ---------- текст ----------

def build_text(conn, monday: date) -> str:
    sunday = monday + timedelta(days=6)
    header = f"📊 Посещаемость за неделю {monday:%d.%m}–{sunday:%d.%m.%Y}"
    totals = week_totals(conn, monday)

    if not totals["complete"]:
        return (f"{header}\n\n⚠️ Похоже, данные за неделю неполные (не отмечены посещения "
                f"или сломалась выгрузка) — проверь CRM.\n{totals['completeness_note']}.")

    stats = week_stats(conn, monday)
    lines = [header, ""]

    change = ""
    if totals["avg_previous"]:
        change = f" ({fmt_pct(totals['visits_total'] / totals['avg_previous'] - 1, signed=True)} к среднему за 4 нед.)"
    lines.append(f"Всего посещений: {totals['visits_total']}{change}")
    lines.append(f"по абонементу {totals['visits_subscription']}, разовых {totals['visits_single']}")
    sessions = f"Занятий состоялось: {totals['sessions_held']}, отменено: {totals['sessions_cancelled']}"
    if totals["sessions_empty"]:
        sessions += f", пустых: {totals['sessions_empty']}"
    lines.append(sessions)
    if totals["holidays"]:
        lines.append(f"Праздники: {', '.join(totals['holidays'])}")

    def block(title, rows, line):
        lines.extend(["", title])
        lines.extend(line(r) for _, r in rows.iterrows()) if len(rows) else lines.append("нет")

    block("🔴 Просели", stats[stats["status"] == STATUS_DROP],
          lambda r: f"• {r['Группа']}\n  {r['status_reason']}, заполняемость {fmt_pct(r['fill_rate'])}")
    block("🟡 Под наблюдением", stats[stats["status"] == STATUS_WATCH],
          lambda r: f"• {r['Группа']}\n  {r['status_reason']}")
    growing = stats[stats["status"] == STATUS_GROW].sort_values("change_pct", ascending=False).head(TOP_GROWING)
    block("🟢 Растут", growing, lambda r: f"• {r['Группа']}\n  {r['status_reason']}")

    changes = week_changes(conn, monday)
    lines.extend(["", "⛔ Отмены и замены"])
    if changes.empty:
        lines.append("нет")
    for _, c in changes.iterrows():
        when = f"{date.fromisoformat(c['date']):%d.%m} {hhmm(c['minutes_begin'])}"
        if c["entity"] == "cancel":
            lines.append(f"• отмена: {c['Группа']}, {when}")
        else:
            lines.append(f"• замена: {c['Группа']}, {when} — вела {c['teacher'] or '?'}")

    new = stats[stats["status"] == STATUS_NEW]
    if len(new):
        lines.extend(["", f"🆕 Новых групп (меньше {config.NEW_GROUP_MIN_WEEKS} нед. истории): {len(new)} — "
                          "цифры в файле, статус пока не ставим"])

    lines.extend(["", "Полная таблица по всем группам — в файле."])
    return "\n".join(lines)


# ---------- Excel ----------

def build_excel(conn, monday: date) -> str:
    stats = week_stats(conn, monday)
    table = pd.DataFrame({
        "Группа": stats["Группа"],
        "Статус": stats["status"],
        "Пояснение": stats["status_reason"],
        "Чел./занятие": stats["avg_per_session"].round(1),
        "Baseline (4 нед.)": stats["baseline"].round(1),
        "Изменение": stats["change_pct"],
        "Посещений": stats["visits_total"],
        "по абонементу": stats["visits_subscription"],
        "разовых": stats["visits_single"],
        "Записались": stats["booked"],
        "Занятий по расписанию": stats["sessions_planned"],
        "Состоялось": stats["sessions_held"],
        "Отменено": stats["sessions_cancelled"],
        "Пустых": stats["sessions_empty"],
        "Вне расписания": stats["sessions_off_schedule"],
        "Заполняемость зала (справочно)": stats["fill_rate"],
    })

    weeks = [w for w in loaded_weeks(conn) if w <= monday][-8:]
    history = pd.read_sql_query(
        f"""SELECT week_start, group_id, avg_per_session FROM weekly_group_stats
            WHERE week_start IN ({','.join('?' * len(weeks))})""",
        conn, params=[w.isoformat() for w in weeks])
    history["Группа"] = history["group_id"].map(group_labels(conn, monday))
    history["Неделя"] = pd.to_datetime(history["week_start"]).dt.strftime("%d.%m")
    dynamics = history.pivot_table(index="Группа", columns="Неделя", values="avg_per_session", aggfunc="first")
    dynamics = dynamics[[w.strftime("%d.%m") for w in weeks if w.strftime("%d.%m") in dynamics.columns]].round(1)

    changes = week_changes(conn, monday)
    changes_out = pd.DataFrame({
        "Дата": pd.to_datetime(changes["date"]).dt.strftime("%d.%m.%Y"),
        "Время": changes["minutes_begin"].map(hhmm),
        "Группа": changes["Группа"],
        "Что": changes["entity"].map({"cancel": "отмена", "substitute": "замена"}),
        "Вела вместо": changes["source_teacher"],
        "Вела": changes["teacher"],
    })

    REPORTS_DIR.mkdir(exist_ok=True)
    path = REPORTS_DIR / f"посещаемость_{monday:%Y-%m-%d}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        table.to_excel(xl, sheet_name="Группы", index=False)
        dynamics.to_excel(xl, sheet_name="Чел. на занятие, 8 нед.")
        changes_out.to_excel(xl, sheet_name="Отмены и замены", index=False)
        for ws in xl.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 9), 60)
            ws.freeze_panes = "B2"
        ws = xl.book["Группы"]
        for row in ws.iter_rows(min_row=2):
            row[5].number_format = "+0%;-0%;0%"
            row[15].number_format = "0%"
    return str(path)


def build_report(conn, monday: date) -> tuple[str, str | None]:
    """Пересчитать показатели и собрать отчёт за неделю. Возвращает (текст, путь к Excel или None)."""
    if monday not in loaded_weeks(conn):
        return (f"За неделю {monday:%d.%m}–{monday + timedelta(days=6):%d.%m.%Y} данных нет: "
                "она ещё не закончилась или не загружена из CRM."), None
    recompute_all(conn)
    return build_text(conn, monday), build_excel(conn, monday)


def parse_date(text: str, today: date | None = None) -> date:
    """«28.09» или «28.09.2026» → дата. Без года — ближайшая прошедшая такая дата."""
    today = today or date.today()
    parts = text.strip().split(".")
    error = ValueError("Дата должна быть в виде ДД.ММ или ДД.ММ.ГГГГ, например 28.09")
    if len(parts) not in (2, 3) or not all(p.isdigit() for p in parts):
        raise error
    day, month = int(parts[0]), int(parts[1])
    year = int(parts[2]) if len(parts) == 3 else today.year
    if len(parts) == 3 and year < 100:
        year += 2000
    try:
        d = date(year, month, day)
    except ValueError:
        raise error from None
    if len(parts) == 2 and d > today:
        d = date(year - 1, month, day)
    return d


def parse_week_arg(text: str, today: date | None = None) -> date:
    """«28.09» → понедельник недели, в которую входит эта дата."""
    return week_of(parse_date(text, today))
